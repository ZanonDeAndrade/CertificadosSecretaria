from __future__ import annotations

import json
from io import BytesIO
from typing import NamedTuple

from certificados_backend.policies import Permission, require_permission
from certificados_backend.schemas.batches import (
    BodyVariablesResponse,
    EmissionBatchPageResponse,
    EmissionBatchResponse,
    SpreadsheetPreviewResponse,
)
from certificados_backend.services import batch_service, job_dispatch, per_work
from certificados_backend.services.certificate_text import BODY_VARIABLES
from certificados_backend.services.spreadsheet import INSTRUCTIONS_SHEET_TITLE
from fastapi import (
    APIRouter,
    BackgroundTasks,
    Depends,
    File,
    Form,
    Header,
    HTTPException,
    Query,
    Request,
    UploadFile,
)
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from database import db
from database.certificate_types import MAX_WORK_AUTHORS, is_per_work_type

router = APIRouter(tags=["emission-batches"])


class _ModelColumn(NamedTuple):
    """A column of the model workbook and how the "Instruções" sheet explains it.

    The example lives only in that sheet: an example left in a data row would
    print on the certificate.
    """

    name: str
    required: bool
    help: str
    example: str
    width: int


_PER_WORK_MODEL: tuple[_ModelColumn, ...] = (
    _ModelColumn("autor1", True, "Nome completo do primeiro autor.", "Maria Oliveira", 28),
    *(
        _ModelColumn(
            f"autor{index}",
            False,
            "Outro autor do mesmo trabalho. Colunas vazias são ignoradas.",
            "João Silva" if index == 2 else "",
            28,
        )
        for index in range(2, MAX_WORK_AUTHORS + 1)
    ),
    _ModelColumn(
        "evento",
        True,
        "Evento em que o trabalho foi apresentado.",
        "Mostra de Iniciação Científica 2026",
        32,
    ),
    _ModelColumn("titulo_trabalho", True, "Título do trabalho.", "Horta vertical automatizada", 36),
    _ModelColumn("equipe", False, "Nome da equipe, se houver.", "Equipe Verde", 24),
    _ModelColumn(
        "carga_horaria",
        False,
        "Horas, só o número. Vazia, o texto automático não cita carga horária.",
        "20",
        16,
    ),
    _ModelColumn(
        "data_emissao",
        False,
        "Data do certificado. Vazia, vale a data informada na tela de emissão.",
        "10/06/2026",
        16,
    ),
    _ModelColumn("data_inicio", False, "Início do evento, se o texto do certificado usar.", "", 16),
    _ModelColumn("data_fim", False, "Fim do evento, se o texto do certificado usar.", "", 16),
)
_ORDINARY_MODEL: tuple[_ModelColumn, ...] = (
    _ModelColumn("nome", True, "Nome completo do participante.", "Maria Oliveira", 32),
    _ModelColumn("carga_horaria", True, "Horas, só o número.", "8", 18),
    _ModelColumn("curso", False, "Curso do participante.", "Sistemas de Informação", 28),
    _ModelColumn("evento", False, "Evento ou atividade.", "Semana Acadêmica 2026", 32),
    _ModelColumn(
        "data_emissao",
        False,
        "Data do certificado. Vazia, vale a data informada na tela de emissão.",
        "10/06/2026",
        18,
    ),
    _ModelColumn("data_inicio", False, "Início do evento, se o texto do certificado usar.", "", 18),
    _ModelColumn("data_fim", False, "Fim do evento, se o texto do certificado usar.", "", 18),
    _ModelColumn("email", False, "E-mail do participante.", "", 32),
    _ModelColumn("documento", False, "Registro Acadêmico (RA) do participante.", "", 20),
    _ModelColumn("equipe", False, "Nome da equipe, se houver.", "", 24),
    _ModelColumn("titulo_trabalho", False, "Título do trabalho, se houver.", "", 36),
)


def _header_style(cells) -> None:
    for cell in cells:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")


@router.get(
    "/certificates/spreadsheet-template",
    operation_id="download_spreadsheet_template",
    responses={
        200: {"content": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {}}}
    },
)
def spreadsheet_template(
    certificate_type: str | None = None,
    _: dict = Depends(require_permission(Permission.CERTIFICATES_ISSUE)),
) -> StreamingResponse:
    """Model workbook for a batch; per-work types get the authors columns.

    The data sheet holds only the header, so nothing in it can reach a
    certificate; the "Instruções" sheet (never read) explains each column.
    """
    per_work_type = is_per_work_type(certificate_type)
    model = _PER_WORK_MODEL if per_work_type else _ORDINARY_MODEL
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "certificados"
    sheet.append([column.name for column in model])
    _header_style(sheet[1])
    for index, column in enumerate(model, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.freeze_panes = "A2"

    guide = workbook.create_sheet(INSTRUCTIONS_SHEET_TITLE)
    guide.append(
        [
            "Preencha a aba “certificados”, uma linha por certificado a partir da linha 2. "
            "Esta aba não é lida na emissão."
        ]
    )
    if per_work_type:
        guide.append(["Um certificado por trabalho, válido para todos os autores."])
    guide.append([])
    guide.append(["Coluna", "Obrigatória", "O que preencher", "Exemplo"])
    _header_style(guide[guide.max_row])
    for column in model:
        guide.append(
            [column.name, "Sim" if column.required else "Não", column.help, column.example]
        )
    for letter, width in zip("ABCD", (18, 12, 72, 36)):
        guide.column_dimensions[letter].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modelo-certificados.xlsx"'},
    )


@router.get(
    "/certificates/body-variables",
    response_model=BodyVariablesResponse,
    operation_id="list_certificate_body_variables",
)
def body_variables(
    _: dict = Depends(require_permission(Permission.CERTIFICATES_ISSUE)),
) -> dict:
    """Variables the body text accepts; the emission screen builds its buttons from it.

    ``per_work`` says whether a per-work type accepts the variable, and
    ``per_work_automatic_text`` is the body such a type prints without a text.
    """
    return {
        "variables": [
            {
                "name": name,
                "token": "{{%s}}" % name,
                "label": label,
                "per_work": name not in per_work.EXCLUDED_BODY_VARIABLES,
            }
            for name, label in BODY_VARIABLES
        ],
        "per_work_automatic_text": per_work.AUTOMATIC_BODY_DESCRIPTION,
    }


@router.post(
    "/certificates/validate-spreadsheet",
    response_model=SpreadsheetPreviewResponse,
    operation_id="validate_certificate_spreadsheet",
)
async def validate_spreadsheet(
    file: UploadFile = File(...),
    data_emissao: str | None = Form(None),
    texto_padrao: str | None = Form(None),
    certificate_type: str | None = Form(None),
    current: dict = Depends(require_permission(Permission.CERTIFICATES_ISSUE)),
) -> dict:
    return await batch_service.preview(
        file,
        data_emissao=data_emissao,
        texto_padrao=texto_padrao,
        certificate_type=certificate_type,
        actor=current,
    )


@router.post(
    "/emission-batches",
    response_model=EmissionBatchResponse,
    status_code=202,
    operation_id="start_emission_batch",
)
async def start_batch(
    request: Request,
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    data_emissao: str | None = Form(None),
    texto_padrao: str | None = Form(None),
    certificate_type: str = Form(...),
    drive_folders_json: str = Form("{}"),
    idempotency_key: str = Header(alias="Idempotency-Key"),
    current: dict = Depends(require_permission(Permission.CERTIFICATES_ISSUE)),
) -> dict:
    try:
        raw_folders = json.loads(drive_folders_json)
    except (TypeError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=422, detail="Destinos do Google Drive inválidos.") from exc
    if not isinstance(raw_folders, dict) or any(
        not isinstance(key, str)
        or not isinstance(folder_id, str)
        or len(key) > 255
        or not (1 <= len(folder_id.strip()) <= 255)
        for key, folder_id in raw_folders.items()
    ):
        raise HTTPException(status_code=422, detail="Destinos do Google Drive inválidos.")
    batch, created = await batch_service.prepare_batch(
        file,
        data_emissao=data_emissao,
        texto_padrao=texto_padrao,
        idempotency_key=idempotency_key,
        drive_folders={key: value.strip() for key, value in raw_folders.items()},
        certificate_type=certificate_type,
        actor=current,
    )
    if created:
        if not int(batch.get("valid_count") or 0):
            batch_service.complete_empty_batch(batch["id"], actor=current)
        elif batch_service.worker_mode() == "inline":
            # Development/test convenience only. The background task receives
            # the durable id, never request-memory rows; a crash leaves a
            # claimable queued/stale job for the external worker.
            background_tasks.add_task(
                batch_service.run_one_job,
                request.app.state.certificate_service,
                batch_id=batch["id"],
            )
        elif job_dispatch.mode() == "cloud_run_job":
            # Queue state was committed before this best-effort wake-up. A
            # periodic fallback execution recovers a failed control-plane call.
            background_tasks.add_task(job_dispatch.dispatch)
    return batch_service.get_batch(batch["id"]) or batch_service.serialize_batch(batch)


def _visible_batch(batch_id: str, current: dict) -> dict:
    batch = batch_service.get_batch(batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Lote não encontrado.")
    return batch


def _mutable_batch(batch_id: str, current: dict) -> dict:
    batch = _visible_batch(batch_id, current)
    if current.get("role") != "admin" and int(batch.get("created_by") or 0) != int(current["id"]):
        raise HTTPException(status_code=403, detail="Operação não permitida.")
    return batch


@router.get(
    "/emission-batches/{batch_id}",
    response_model=EmissionBatchResponse,
    operation_id="get_emission_batch",
)
def get_batch(
    batch_id: str,
    current: dict = Depends(require_permission(Permission.CERTIFICATES_READ)),
) -> dict:
    return _visible_batch(batch_id, current)


@router.post(
    "/emission-batches/{batch_id}/cancel",
    response_model=EmissionBatchResponse,
    operation_id="cancel_emission_batch",
)
def cancel_batch(
    batch_id: str,
    current: dict = Depends(require_permission(Permission.CERTIFICATES_ISSUE)),
) -> dict:
    _mutable_batch(batch_id, current)
    return batch_service.cancel_batch(batch_id, actor=current)


@router.post(
    "/emission-batches/{batch_id}/retry",
    response_model=EmissionBatchResponse,
    operation_id="retry_emission_batch",
)
def retry_batch(
    request: Request,
    background_tasks: BackgroundTasks,
    batch_id: str,
    current: dict = Depends(require_permission(Permission.CERTIFICATES_ISSUE)),
) -> dict:
    _mutable_batch(batch_id, current)
    retried = batch_service.retry_batch(batch_id, actor=current)
    if batch_service.worker_mode() == "inline":
        background_tasks.add_task(
            batch_service.run_one_job,
            request.app.state.certificate_service,
            batch_id=batch_id,
        )
    elif job_dispatch.mode() == "cloud_run_job":
        background_tasks.add_task(job_dispatch.dispatch)
    return retried


@router.get(
    "/emission-batches/by-idempotency-key/{idempotency_key}",
    response_model=EmissionBatchResponse,
    operation_id="resume_emission_batch",
)
def by_idempotency_key(
    idempotency_key: str,
    current: dict = Depends(require_permission(Permission.CERTIFICATES_ISSUE)),
) -> dict:
    key = batch_service.validate_idempotency_key(idempotency_key)
    row = db.get_emission_batch_by_idempotency_key(key)
    if not row or (
        current.get("role") != "admin" and int(row.get("created_by") or 0) != int(current["id"])
    ):
        raise HTTPException(status_code=404, detail="Lote não encontrado.")
    return batch_service.with_work_details(batch_service.serialize_batch(row))


@router.get(
    "/emission-batches",
    response_model=EmissionBatchPageResponse,
    operation_id="list_emission_batches",
)
def list_batches(
    limit: int = Query(default=20, ge=1, le=100),
    offset: int = Query(default=0, ge=0),
    current: dict = Depends(require_permission(Permission.CERTIFICATES_READ)),
) -> dict:
    rows, total = db.list_emission_batches(limit=limit, offset=offset)
    return {
        "items": [batch_service.serialize_batch(row) for row in rows],
        "total": total,
        "limit": limit,
        "offset": offset,
    }
