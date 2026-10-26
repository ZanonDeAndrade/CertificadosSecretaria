"""Per-work batches never fall back to the generic template.

A per-work type (the Mostra de Iniciação Científica) is refused unless its own
active template has the authors block, a body and no participant name. Its
certificate highlights no one: the body names the work and the event and the
block lists the authors under "Autores:" (or "Autor:"). Blank text means the
automatic body. The spreadsheet requires the work title and the event, and a
cell still holding an example text of the model workbook invalidates its row.
Ordinary types keep their template, text and PDF exactly as before.
"""

from __future__ import annotations

import base64
import io
from io import BytesIO
from itertools import count

import pandas as pd
import pytest
from certificados_backend.services import batch_service, per_work, template_service
from certificados_backend.services.spreadsheet import (
    TEMPLATE_EXAMPLE_TEXTS,
    SpreadsheetError,
    SpreadsheetRow,
    read_and_validate,
)
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw

from database import db

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MOSTRA = "scientific_initiation"
EVENT = "Mostra de Iniciação Científica 2026"
TITLE = "Horta vertical automatizada"
PREAMBLE = "Certificamos, para os devidos fins, que"
AUTHORS = (
    "Pessoa Exemplo Alfa",
    "Pessoa Exemplo Beta",
    "Pessoa Exemplo Gama",
    "Pessoa Exemplo Delta",
    "Pessoa Exemplo Épsilon",
    "Pessoa Exemplo Zeta",
    "Pessoa Exemplo Eta",
    "Pessoa Exemplo Teta",
)
ORDINARY_BODY = (
    "participou da atividade descrita neste certificado, com carga horária total de "
    "{{carga_horaria}} horas."
)
_KEYS = count(1)


def _work(*authors: str, **extra) -> dict:
    """A row with the columns of the Mostra model workbook."""
    row: dict = {f"autor{index}": "" for index in range(1, 9)}
    row.update({f"autor{index}": name for index, name in enumerate(authors, start=1)})
    row.update(
        evento=EVENT,
        titulo_trabalho=TITLE,
        equipe="",
        carga_horaria="",
        data_emissao="10/06/2026",
        data_inicio="",
        data_fim="",
    )
    row.update(extra)
    return row


def _xlsx(rows: list[dict]) -> bytes:
    buffer = io.BytesIO()
    pd.DataFrame(rows).to_excel(buffer, index=False)
    return buffer.getvalue()


def _files(rows: list[dict]) -> dict:
    return {"file": ("lote.xlsx", _xlsx(rows), XLSX_MIME)}


def _preview(client, rows: list[dict], *, text: str = "", certificate_type: str = MOSTRA):
    return client.post(
        "/certificates/validate-spreadsheet",
        files=_files(rows),
        data={"texto_padrao": text, "certificate_type": certificate_type},
    )


def _start(client, rows: list[dict], *, text: str = "", certificate_type: str = MOSTRA):
    return client.post(
        "/emission-batches",
        files=_files(rows),
        data={
            "texto_padrao": text,
            "certificate_type": certificate_type,
            "drive_folders_json": "{}",
        },
        headers={"Idempotency-Key": f"per-work-test-{next(_KEYS):06d}"},
    )


def _record_texts(monkeypatch) -> list[str]:
    texts: list[str] = []
    original = ImageDraw.ImageDraw.text

    def recording(self, xy, text, fill=None, font=None, *args, **kwargs):
        texts.append(text)
        return original(self, xy, text, fill, font, *args, **kwargs)

    monkeypatch.setattr(ImageDraw.ImageDraw, "text", recording)
    return texts


def _mostra_version(elements: list[dict]) -> dict:
    buffer = BytesIO()
    Image.new("RGB", (1414, 1000), "white").save(buffer, "PNG")
    layout = {
        "background": "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode(),
        "image_width": 1414,
        "image_height": 1000,
        "elements": elements,
    }
    version = template_service.create_version(name="Mostra", layout=layout, certificate_type=MOSTRA)
    assert template_service.activate_version(version["id"])
    return version


def _text(key: str, **extra) -> dict:
    return {"type": "text", "key": key, "x": 70, "y": 300, "fontSize": 24, **extra}


QR = {"type": "qr", "key": "qr", "x": 70, "y": 820, "width": 120, "height": 120}


# ── No silent fallback to the generic template ─────────────────────────────────


def test_mostra_batch_without_a_template_of_its_type_is_refused(client, auth_headers):
    # The generic template is active, as in production: it must not be used.
    assert template_service.get_active_version(None) is not None
    assert template_service.get_active_version(MOSTRA) is None
    rows = [_work(*AUTHORS[:3], carga_horaria=60)]

    preview = _preview(client, rows)
    assert preview.status_code == 200, preview.text
    problem = preview.json()["template_problem"]
    assert problem.startswith("Não há template ativo do tipo Mostra de Iniciação Científica.")
    assert "Em Template" in problem and "ative" in problem
    assert preview.json()["can_manage_templates"] is True

    started = _start(client, rows)
    assert started.status_code == 422
    assert started.json()["detail"] == problem
    assert db.list_emission_batches(limit=10, offset=0)[1] == 0
    assert client.get("/certificates").json()["total"] == 0


@pytest.mark.parametrize(
    ("elements", "fragment"),
    [
        (
            [_text("texto_certificado"), QR],
            "não tem o bloco Autores. Em Template, use “Editar como nova” nessa versão, "
            "adicione o bloco Autores acima do QR Code, salve e ative a nova versão.",
        ),
        (
            [_text("name"), _text("texto_certificado"), _text("authors", anchorY="bottom"), QR],
            "tem o elemento Nome do participante. Em Template, use “Editar como nova” "
            "nessa versão, remova o elemento Nome, salve e ative a nova versão.",
        ),
        (
            [_text("name"), _text("authors", anchorY="bottom"), QR],
            "tem o elemento Nome do participante e não tem o elemento Corpo.",
        ),
    ],
    ids=["sem-bloco-autores", "com-nome-do-participante", "nome-e-sem-corpo"],
)
def test_mostra_template_without_authors_or_with_the_name_is_refused(
    client, auth_headers, elements, fragment
):
    _mostra_version(elements)
    rows = [_work(*AUTHORS[:3])]

    preview = _preview(client, rows)
    assert preview.status_code == 200, preview.text
    assert fragment in preview.json()["template_problem"]

    started = _start(client, rows)
    assert started.status_code == 422
    assert fragment in started.json()["detail"]
    assert client.get("/certificates").json()["total"] == 0


def test_refusal_tells_who_cannot_manage_templates_to_ask_an_administrator():
    missing = per_work.template_problem(None, MOSTRA, can_manage=False)
    assert missing is not None and "Peça a um administrador" in missing
    unfit = per_work.template_problem(
        {"layout": {"elements": [_text("name"), _text("texto_certificado")]}},
        MOSTRA,
        can_manage=False,
    )
    assert unfit is not None and "Peça a um administrador" in unfit
    # Ordinary types are never checked.
    assert per_work.template_problem(None, "listener", can_manage=True) is None


def test_worker_issues_nothing_when_the_template_stops_fitting_after_confirmation(
    client, auth_headers, mostra_template, monkeypatch
):
    monkeypatch.setenv("BATCH_WORKER_MODE", "external")
    started = _start(client, [_work(*AUTHORS[:3])])
    assert started.status_code == 202, started.text
    batch_id = started.json()["id"]

    # A Mostra version with the participant name becomes active meanwhile.
    _mostra_version([_text("name"), _text("texto_certificado"), _text("authors"), QR])
    assert batch_service.run_one_job(client.app.state.certificate_service, batch_id=batch_id)

    assert client.get("/certificates").json()["total"] == 0
    assert db.get_emission_batch(batch_id)["status"] != "completed"


# ── The Mostra certificate ─────────────────────────────────────────────────────


def test_mostra_pdf_highlights_no_name_and_lists_one_three_and_eight_authors(
    client, auth_headers, mostra_template, monkeypatch
):
    texts = _record_texts(monkeypatch)
    rows = [
        _work(AUTHORS[0], titulo_trabalho="Trabalho com um autor"),
        _work(*AUTHORS[:3], carga_horaria=60),
        _work(*AUTHORS, titulo_trabalho="Trabalho com oito autores"),
    ]
    preview = _preview(client, rows)
    assert preview.status_code == 200, preview.text
    assert preview.json()["template_problem"] is None
    assert preview.json()["resolved_text_preview"] == (
        f'Certificamos que o trabalho "Trabalho com um autor" foi apresentado no evento {EVENT}.'
    )

    started = _start(client, rows)
    assert started.status_code == 202, started.text
    batch = client.get(f"/emission-batches/{started.json()['id']}").json()
    assert batch["status"] == "completed" and batch["generated_count"] == 3

    blocks = [text for text in texts if text.startswith(("Autor:", "Autores:"))]
    assert blocks == [
        f"Autor:\n{AUTHORS[0]}",
        "Autores:\n" + "\n".join(AUTHORS[:3]),
        "Autores:\n" + "\n".join(AUTHORS),
    ]
    # No highlighted name and no preamble announcing one.
    assert not set(texts) & set(AUTHORS)
    assert PREAMBLE not in texts
    assert not any(text.startswith("Autor 1") for text in texts)

    drawn = " ".join(texts)
    # The workload is cited only when the spreadsheet brings it.
    assert (
        f'Certificamos que o trabalho "{TITLE}" foi apresentado no evento {EVENT}, '
        "com carga horária de 60 horas." in drawn
    )
    assert (
        'Certificamos que o trabalho "Trabalho com oito autores" foi apresentado no evento '
        f"{EVENT}." in drawn
    )

    items = {item["work_title"]: item for item in client.get("/certificates").json()["items"]}
    assert items["Trabalho com oito autores"]["authors"] == list(AUTHORS)
    stored = db.get_by_code(items[TITLE]["unique_code"])
    assert stored["certificate_text"] == (
        f'Certificamos que o trabalho "{TITLE}" foi apresentado no evento {EVENT}, '
        "com carga horária de 60 horas."
    )


def test_a_written_text_replaces_the_automatic_one_but_cannot_use_nome(
    client, auth_headers, mostra_template
):
    rows = [_work(*AUTHORS[:2])]
    own = _preview(client, rows, text="Certificamos que {{titulo_trabalho}} foi apresentado.")
    assert own.status_code == 200, own.text
    assert own.json()["resolved_text_preview"] == f"Certificamos que {TITLE} foi apresentado."

    with_name = _preview(client, rows, text="{{nome}} apresentou {{titulo_trabalho}}.")
    assert with_name.status_code == 400
    assert "não pode usar {{nome}}" in with_name.json()["detail"]
    assert _start(client, rows, text="{{ nome }} apresentou.").status_code == 400

    # Whitespace only is no text at all: the automatic body applies.
    blank = _preview(client, rows, text="  \n ")
    assert blank.status_code == 200
    assert blank.json()["resolved_text_preview"].startswith('Certificamos que o trabalho "')


def test_automatic_body_names_the_work_and_the_event():
    row = SpreadsheetRow(
        row_number=2, nome="x", carga_horaria=None, evento="Feira X", titulo_trabalho="Robô"
    )
    assert per_work.automatic_body(row) == (
        'Certificamos que o trabalho "Robô" foi apresentado no evento Feira X.'
    )
    with_hours = SpreadsheetRow(
        row_number=2, nome="x", carga_horaria=8, evento="Feira X", titulo_trabalho="Robô"
    )
    assert per_work.automatic_body(with_hours).endswith(", com carga horária de 8 horas.")


def test_body_variables_say_which_ones_a_per_work_type_accepts(client, auth_headers):
    response = client.get("/certificates/body-variables")
    assert response.status_code == 200
    accepted = {item["name"]: item["per_work"] for item in response.json()["variables"]}
    assert accepted["nome"] is False
    assert accepted["titulo_trabalho"] is True and accepted["evento"] is True
    assert response.json()["per_work_automatic_text"] == per_work.AUTOMATIC_BODY_DESCRIPTION


# ── Ordinary types are unchanged ──────────────────────────────────────────────


def test_ordinary_batches_keep_the_generic_template_text_and_name(
    client, auth_headers, mostra_template, monkeypatch
):
    texts = _record_texts(monkeypatch)
    row = {"nome": "Ana Souza", "evento": "Semana Jurídica", "carga_horaria": 40}

    assert _preview(client, [row], text="", certificate_type="listener").status_code == 400
    preview = _preview(client, [row], text=ORDINARY_BODY, certificate_type="listener")
    assert preview.status_code == 200
    assert preview.json()["template_problem"] is None
    # A dedicated Mostra template changes nothing for the other types.
    assert template_service.get_active_version("listener")["certificate_type"] is None

    started = _start(client, [row], text=ORDINARY_BODY, certificate_type="listener")
    assert started.status_code == 202, started.text
    assert PREAMBLE in texts and "Ana Souza" in texts
    assert (
        "participou da atividade descrita neste certificado, com carga horária total de "
        "40 horas." in " ".join(texts)
    )
    assert not any(text.startswith(("Autor:", "Autores:")) for text in texts)


# ── Spreadsheet ────────────────────────────────────────────────────────────────


def test_mostra_rows_need_the_work_title_and_the_event():
    report = read_and_validate(
        BytesIO(
            _xlsx(
                [
                    _work(AUTHORS[0], titulo_trabalho=""),
                    _work(AUTHORS[1], evento=""),
                    _work(AUTHORS[2]),
                ]
            )
        ),
        certificate_type=MOSTRA,
    )
    assert [(row.row_number, row.errors) for row in report.invalid] == [
        (2, ["titulo_trabalho é obrigatório"]),
        (3, ["evento é obrigatório"]),
    ]
    assert [row.nome for row in report.valid] == [AUTHORS[2]]

    without_title = [{"autor1": AUTHORS[0], "evento": EVENT}]
    with pytest.raises(SpreadsheetError, match="ausentes na planilha: titulo_trabalho$"):
        read_and_validate(BytesIO(_xlsx(without_title)), certificate_type=MOSTRA)


@pytest.mark.parametrize("certificate_type", [MOSTRA, "listener"])
def test_rows_left_with_an_example_text_of_the_model_are_invalid(certificate_type):
    if certificate_type == MOSTRA:
        old_model_row = _work(
            "Nome do primeiro autor",
            "Segundo autor (opcional)",
            evento="Nome da mostra",
            titulo_trabalho="Título do trabalho (opcional)",
            equipe="Equipe (opcional)",
        )
        filled = _work(*AUTHORS[:2], equipe=" equipe  (OPCIONAL) ")
        good = _work(*AUTHORS[:2])
    else:
        old_model_row = {
            "nome": "Nome completo",
            "carga_horaria": 8,
            "curso": "Curso (opcional)",
            "evento": "Evento (opcional)",
            "equipe": "Equipe (opcional)",
            "titulo_trabalho": "Título do trabalho (opcional)",
        }
        filled = {"nome": "Ana Souza", "carga_horaria": 8, "equipe": " equipe  (OPCIONAL) "}
        good = {"nome": "Ana Souza", "carga_horaria": 8, "equipe": "Equipe Azul"}
    filled_columns = {key: "" for key in old_model_row} | filled
    good_columns = {key: "" for key in old_model_row} | good
    report = read_and_validate(
        BytesIO(_xlsx([old_model_row, filled_columns, good_columns])),
        certificate_type=certificate_type,
    )

    assert [row.row_number for row in report.invalid] == [2, 3]
    old_errors = report.invalid[0].errors
    assert len(old_errors) == len(
        [value for value in old_model_row.values() if value in TEMPLATE_EXAMPLE_TEXTS]
    )
    assert any(
        error == 'equipe contém o texto de exemplo do modelo ("Equipe (opcional)"); '
        "substitua pelo dado real ou apague"
        for error in old_errors
    )
    # Case and spacing do not hide the example text.
    assert report.invalid[1].errors == [
        'equipe contém o texto de exemplo do modelo ("equipe  (OPCIONAL)"); '
        "substitua pelo dado real ou apague"
    ]
    assert len(report.valid) == 1


@pytest.mark.parametrize("certificate_type", [MOSTRA, None])
def test_model_workbook_has_no_example_row_and_explains_each_column(
    client, auth_headers, certificate_type
):
    params = {"certificate_type": certificate_type} if certificate_type else {}
    response = client.get("/certificates/spreadsheet-template", params=params)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["certificados", "Instruções"]
    data = workbook["certificados"]
    headers = [cell.value for cell in data[1]]
    # Only the header: nothing in the data sheet can reach a certificate.
    assert data.max_row == 1

    guide = list(workbook["Instruções"].iter_rows(values_only=True))
    table = guide[guide.index(("Coluna", "Obrigatória", "O que preencher", "Exemplo")) + 1 :]
    assert [line[0] for line in table] == headers
    required = {line[0] for line in table if line[1] == "Sim"}
    if certificate_type == MOSTRA:
        assert required == {"autor1", "evento", "titulo_trabalho"}
    else:
        assert required == {"nome", "carga_horaria"}
    examples = {str(line[3]) for line in table if line[3]}
    assert not examples & set(TEMPLATE_EXAMPLE_TEXTS)


def test_filled_model_is_read_from_its_data_sheet_even_with_instructions_active(
    client, auth_headers
):
    response = client.get("/certificates/spreadsheet-template", params={"certificate_type": MOSTRA})
    workbook = load_workbook(BytesIO(response.content))
    data = workbook["certificados"]
    headers = [cell.value for cell in data[1]]
    values = _work(*AUTHORS[:3], carga_horaria=60)
    data.append([values.get(header, "") for header in headers])
    # Saved with the instructions sheet open, as a user may leave it.
    workbook.active = workbook.sheetnames.index("Instruções")
    buffer = BytesIO()
    workbook.save(buffer)

    report = read_and_validate(BytesIO(buffer.getvalue()), certificate_type=MOSTRA)
    assert report.invalid == []
    assert [row.autores for row in report.valid] == [AUTHORS[:3]]
    assert report.valid[0].titulo_trabalho == TITLE


def test_a_workbook_without_an_instructions_sheet_still_reads_its_first_sheet():
    workbook = Workbook()
    first = workbook.active
    first.append(["nome", "carga_horaria"])
    first.append(["Ana Souza", 8])
    other = workbook.create_sheet("Rascunho")
    other.append(["nome", "carga_horaria"])
    other.append(["Outra Pessoa", 4])
    workbook.active = 1
    buffer = BytesIO()
    workbook.save(buffer)

    report = read_and_validate(BytesIO(buffer.getvalue()))
    assert [row.nome for row in report.valid] == ["Ana Souza"]
