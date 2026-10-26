"""Per-work certificates (Mostra de Iniciação Científica): author columns, the
authors block that grows upward from a fixed base, {{equipe}} and
{{titulo_trabalho}}, templates dedicated to a type, public validation listing
every author, and legacy templates rendering exactly as before."""

from __future__ import annotations

import base64
import importlib.util
import io
import sqlite3
import uuid
from io import BytesIO
from itertools import count
from pathlib import Path

import pandas as pd
import pytest
from certificados_backend.models import ParticipantRegistryRecord
from certificados_backend.services import template_service
from certificados_backend.services.certificate_text import (
    render_certificate_body,
    validate_body_template,
)
from certificados_backend.services.generator import (
    CertificateGenerator,
    CertificateGeneratorConfig,
    authors_block,
    measuring_draw,
)
from certificados_backend.services.spreadsheet import (
    SpreadsheetError,
    compute_business_key,
    read_and_validate,
)
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw

from database import db

_BACKEND_DIR = Path(__file__).resolve().parent.parent
_REPO_ROOT = _BACKEND_DIR.parent.parent
IMAGE_W, IMAGE_H = 1414, 1000
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MOSTRA = "scientific_initiation"
BODY = "apresentou o trabalho {{titulo_trabalho}} pela equipe {{equipe}}."
_KEYS = count(1)
AUTHORS = (
    "Pessoa Exemplo Alfa",
    "Pessoa Exemplo Beta",
    "Pessoa Exemplo Gama",
    "Pessoa Exemplo Delta",
    "Pessoa Exemplo Épsilon",
    "Pessoa Exemplo Zeta",
    "Pessoa Exemplo Eta",
    "Pessoa Exemplo Teta",
)


def _xlsx(rows: list[dict]) -> bytes:
    buffer = io.BytesIO()
    pd.DataFrame(rows).to_excel(buffer, index=False)
    return buffer.getvalue()


def _files(rows: list[dict]) -> dict:
    return {"file": ("mostra.xlsx", _xlsx(rows), XLSX_MIME)}


def _work(*authors: str, **extra) -> dict:
    row: dict = {f"autor{index}": name for index, name in enumerate(authors, start=1)}
    row.setdefault("evento", "Mostra de Iniciação Científica 2026")
    row.setdefault("titulo_trabalho", "Horta vertical automatizada")
    row.setdefault("equipe", "Equipe Verde")
    row.setdefault("data_emissao", "10/06/2026")
    row.update(extra)
    return row


def _png_bytes(w: int = IMAGE_W, h: int = IMAGE_H) -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (w, h), "white").save(buffer, "PNG")
    return buffer.getvalue()


def _generator(tmp_path: Path) -> CertificateGenerator:
    return CertificateGenerator(
        CertificateGeneratorConfig(
            template_path=_BACKEND_DIR / "templates" / "certificado_base.png",
            regular_font_path=_BACKEND_DIR / "fonts" / "times.ttf",
            bold_font_path=_BACKEND_DIR / "fonts" / "timesbd.ttf",
            output_dir=tmp_path / "pdfs",
        )
    )


def _record(authors: tuple[str, ...]) -> ParticipantRegistryRecord:
    return ParticipantRegistryRecord(
        nome=authors[0] if authors else "Pessoa Exemplo Alfa",
        email="",
        curso="",
        livro=0,
        folha=0,
        linha=0,
        evento="Mostra 2026",
        validation_code="CERT-2026-WA0001",
        texto_certificado="apresentou o trabalho.",
        certificate_text="apresentou o trabalho.",
        data_emissao="10 de junho de 2026",
        certificate_type="Mostra de Iniciação Científica",
        autores=authors,
        equipe="Equipe Verde",
        titulo_trabalho="Horta vertical automatizada",
    )


def _record_text_calls(monkeypatch) -> list[dict]:
    calls: list[dict] = []
    original = ImageDraw.ImageDraw.text

    def recording(self, xy, text, fill=None, font=None, *args, **kwargs):
        calls.append({"xy": xy, "text": text})
        return original(self, xy, text, fill, font, *args, **kwargs)

    monkeypatch.setattr(ImageDraw.ImageDraw, "text", recording)
    return calls


def _layout(elements: list[dict]) -> dict:
    return {
        "background": "data:image/png;base64," + base64.b64encode(_png_bytes()).decode("ascii"),
        "image_width": IMAGE_W,
        "image_height": IMAGE_H,
        "elements": elements,
    }


def _text(key: str, **overrides) -> dict:
    element = {
        "type": "text",
        "key": key,
        "x": 90,
        "y": 300,
        "fontSize": 24,
        "fontFamily": "Times New Roman",
        "color": "#000000",
        "align": "left",
        "bold": False,
        "italic": False,
    }
    element.update(overrides)
    return element


# Fit for a per-work type: no participant name, a body, the authors block.
MOSTRA_ELEMENTS = [
    _text("texto_certificado", y=320, width=1000),
    _text("authors", y=900, anchorY="bottom"),
    {"type": "qr", "key": "qr", "x": 1200, "y": 820, "width": 140, "height": 140},
]


# ── Spreadsheet ───────────────────────────────────────────────────────────────


def test_reads_one_three_and_eight_authors_and_compacts_blank_columns():
    rows = [
        _work(AUTHORS[0]),
        _work(*AUTHORS[:3]),
        _work(*AUTHORS, carga_horaria=20),
        # Blank columns between authors are dropped, order is kept.
        {**_work(AUTHORS[0]), "autor2": "", "autor3": AUTHORS[2], "autor5": AUTHORS[4]},
    ]
    report = read_and_validate(BytesIO(_xlsx(rows)), certificate_type=MOSTRA)

    assert report.invalid == []
    assert [len(row.autores) for row in report.valid] == [1, 3, 8, 3]
    assert report.valid[3].autores == (AUTHORS[0], AUTHORS[2], AUTHORS[4])
    # The first author is the holder; the workload may be blank.
    assert [row.nome for row in report.valid] == [AUTHORS[0]] * 4
    assert report.valid[0].carga_horaria is None
    assert report.valid[2].carga_horaria == 20
    assert report.valid[0].equipe == "Equipe Verde"
    assert report.valid[0].titulo_trabalho == "Horta vertical automatizada"


def test_autor9_column_is_ignored_with_a_warning():
    rows = [{**_work(*AUTHORS), "Autor 9": "Pessoa Exemplo Iota", "autor10": "Pessoa Kapa"}]
    report = read_and_validate(BytesIO(_xlsx(rows)), certificate_type=MOSTRA)

    assert report.valid[0].autores == AUTHORS
    assert len(report.warnings) == 1
    assert "Autor 9" in report.warnings[0] and "autor10" in report.warnings[0]
    assert "8 autores" in report.warnings[0]


def test_header_variations_map_to_the_author_columns():
    rows = [
        {
            "Autor 1": AUTHORS[0],
            "autor_2": AUTHORS[1],
            "AUTOR03": AUTHORS[2],
            "Evento": "Mostra",
            "Título do trabalho": "Horta",
        }
    ]
    report = read_and_validate(BytesIO(_xlsx(rows)), certificate_type=MOSTRA)
    assert report.valid[0].autores == AUTHORS[:3]


def test_nome_stands_in_for_autor1_only_when_autor1_is_absent():
    work = {"evento": "Mostra", "titulo_trabalho": "Horta"}
    only_name = read_and_validate(
        BytesIO(_xlsx([{"nome": AUTHORS[0], **work}])), certificate_type=MOSTRA
    )
    assert only_name.valid[0].autores == (AUTHORS[0],)
    assert only_name.valid[0].nome == AUTHORS[0]

    both = read_and_validate(
        BytesIO(_xlsx([{"nome": "Ignorado", "autor1": AUTHORS[0], **work}])),
        certificate_type=MOSTRA,
    )
    assert both.valid[0].nome == AUTHORS[0]

    with pytest.raises(SpreadsheetError, match="ausentes na planilha: autor1$"):
        read_and_validate(BytesIO(_xlsx([work])), certificate_type=MOSTRA)


def test_work_without_first_author_is_invalid_and_other_types_are_unchanged():
    report = read_and_validate(
        BytesIO(_xlsx([_work("", carga_horaria=8), _work(AUTHORS[0])])), certificate_type=MOSTRA
    )
    assert [row.errors for row in report.invalid] == [["autor1 é obrigatório"]]
    assert len(report.valid) == 1

    # Ordinary types keep requiring nome and carga_horaria.
    ordinary = read_and_validate(
        BytesIO(_xlsx([{"nome": "Ana", "carga_horaria": ""}])), certificate_type="listener"
    )
    assert ordinary.invalid[0].errors == ["carga_horaria é obrigatória"]
    with pytest.raises(SpreadsheetError, match="nome"):
        read_and_validate(BytesIO(_xlsx([{"autor1": "Ana", "carga_horaria": 8}])))


def test_business_key_separates_works_that_share_the_first_author():
    report = read_and_validate(
        BytesIO(
            _xlsx(
                [
                    _work(AUTHORS[0], AUTHORS[1], titulo_trabalho="Trabalho A"),
                    _work(AUTHORS[0], AUTHORS[2], titulo_trabalho="Trabalho B"),
                ]
            )
        ),
        certificate_type=MOSTRA,
    )
    first, second = report.valid
    assert compute_business_key(first, BODY, MOSTRA) != compute_business_key(second, BODY, MOSTRA)


# ── Rendering ─────────────────────────────────────────────────────────────────


def test_authors_block_grows_upward_from_a_fixed_base(tmp_path, monkeypatch):
    generator = _generator(tmp_path)
    draw = measuring_draw()
    element = _text("authors", y=900, anchorY="bottom")

    three = generator.layout_text_box(draw, IMAGE_W, authors_block(list(AUTHORS[:3])), element)
    eight = generator.layout_text_box(draw, IMAGE_W, authors_block(list(AUTHORS)), element)
    assert eight.text_height > three.text_height

    layout = {"image_width": IMAGE_W, "image_height": IMAGE_H, "elements": [element]}
    calls = _record_text_calls(monkeypatch)
    generator._compose_visual_image(_record(AUTHORS[:3]), layout, background_bytes=_png_bytes())
    generator._compose_visual_image(_record(AUTHORS), layout, background_bytes=_png_bytes())

    assert calls[0]["text"] == "Autores:\n" + "\n".join(AUTHORS[:3])
    # Heading plus eight authors, unnumbered.
    assert calls[1]["text"].split("\n") == ["Autores:", *AUTHORS]
    # The base stays at Y = 900; more authors push the first line further up.
    assert calls[0]["xy"][1] == 900 - three.text_height
    assert calls[1]["xy"][1] == 900 - eight.text_height
    assert calls[1]["xy"][1] < calls[0]["xy"][1]


def test_participant_data_exposes_authors_team_and_title(tmp_path):
    generator = _generator(tmp_path)
    data = generator.participant_data(_record(AUTHORS[:2]))
    assert data["autor1"] == AUTHORS[0] and data["autor2"] == AUTHORS[1]
    assert data["autor3"] == "" and data["autor8"] == ""
    assert data["authors"] == f"Autores:\n{AUTHORS[0]}\n{AUTHORS[1]}"
    # A single author gets the singular heading.
    assert generator.participant_data(_record(AUTHORS[:1]))["authors"] == f"Autor:\n{AUTHORS[0]}"
    assert data["equipe"] == "Equipe Verde"
    assert data["work_title"] == "Horta vertical automatizada"

    # Ordinary certificates carry none of it, so their elements stay empty.
    plain = generator.participant_data(_record(()))
    assert plain["authors"] == "" and plain["autor1"] == ""


def test_equipe_and_titulo_trabalho_are_body_variables():
    assert validate_body_template(BODY) == BODY
    rendered = render_certificate_body(
        BODY, {"titulo_trabalho": "Horta vertical", "equipe": "Equipe Verde"}
    )
    assert rendered == "apresentou o trabalho Horta vertical pela equipe Equipe Verde."


# ── Templates per type ────────────────────────────────────────────────────────


def _create(client, name: str, certificate_type: str | None, elements=None) -> dict:
    response = client.post(
        "/templates/versions",
        json={
            "name": name,
            "layout": _layout(elements or [_text("name")]),
            "certificate_type": certificate_type,
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_each_type_activates_its_own_version_and_falls_back_to_the_generic(client, auth_headers):
    generic = template_service.get_active_version()
    assert generic and generic["certificate_type"] is None

    mostra_v1 = _create(client, "Mostra v1", MOSTRA)
    assert mostra_v1["certificate_type"] == MOSTRA
    assert client.post(f"/templates/versions/{mostra_v1['id']}/activate").status_code == 200

    # The generic version is still active; the Mostra has its own.
    assert template_service.get_active_version()["id"] == generic["id"]
    assert template_service.get_active_version(MOSTRA)["id"] == mostra_v1["id"]
    assert template_service.get_active_version("listener")["id"] == generic["id"]
    by_api = client.get("/templates/active", params={"certificate_type": MOSTRA})
    assert by_api.status_code == 200 and by_api.json()["id"] == mostra_v1["id"]

    # Activating another Mostra version steps down only the previous Mostra one.
    mostra_v2 = _create(client, "Mostra v2", MOSTRA)
    assert client.post(f"/templates/versions/{mostra_v2['id']}/activate").status_code == 200
    versions = {item["id"]: item for item in client.get("/templates/versions").json()}
    assert versions[mostra_v1["id"]]["is_active"] is False
    assert versions[mostra_v2["id"]]["is_active"] is True
    assert versions[generic["id"]]["is_active"] is True

    unknown = client.post(
        "/templates/versions",
        json={"name": "x", "layout": _layout([_text("name")]), "certificate_type": "nope"},
    )
    assert unknown.status_code == 400


def test_database_refuses_two_active_versions_for_the_same_type(client, auth_headers):
    """The unique index on ``active_type_key`` is the guarantee, not the code."""
    second = _create(client, "Outra genérica", None)
    connection = sqlite3.connect(str(db.DATABASE_PATH))
    try:
        with pytest.raises(sqlite3.IntegrityError):
            connection.execute(
                "UPDATE template_versions SET is_active = 1, active_type_key = '__generic__' "
                "WHERE id = ?",
                (second["id"],),
            )
    finally:
        connection.close()


# ── End to end: batch, history, public validation ─────────────────────────────


def _public_client(monkeypatch) -> TestClient:
    monkeypatch.setenv("DOCUMENT_HASH_SECRET", "T8v!mZ2#qL9@pR4$xN7&cB1*eK6-wY3+")
    monkeypatch.setenv("DOCUMENT_HASH_ACTIVE_KEY_ID", "v1")
    path = _REPO_ROOT / "certificados_consulta" / "app.py"
    spec = importlib.util.spec_from_file_location(f"consulta_{uuid.uuid4().hex}", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return TestClient(module.app)


def test_mostra_batch_issues_one_certificate_per_work_and_lists_the_authors(
    client, auth_headers, monkeypatch
):
    mostra = _create(client, "Mostra", MOSTRA, MOSTRA_ELEMENTS)
    assert client.post(f"/templates/versions/{mostra['id']}/activate").status_code == 200

    preview = client.post(
        "/certificates/validate-spreadsheet",
        files=_files([_work(*AUTHORS[:3]), _work(AUTHORS[3])]),
        data={"texto_padrao": BODY, "certificate_type": MOSTRA},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["valid"][0]["autores"] == list(AUTHORS[:3])
    assert preview.json()["valid"][0]["carga_horaria"] is None
    assert preview.json()["missing_settings"] == []

    started = client.post(
        "/emission-batches",
        files=_files([_work(*AUTHORS[:3]), _work(AUTHORS[3])]),
        data={"texto_padrao": BODY, "certificate_type": MOSTRA, "drive_folders_json": "{}"},
        headers={"Idempotency-Key": f"work-authors-test-{next(_KEYS):06d}"},
    )
    assert started.status_code == 202, started.text
    batch = client.get(f"/emission-batches/{started.json()['id']}").json()
    assert batch["status"] == "completed", batch
    assert batch["generated_count"] == 2

    items = client.get("/certificates").json()["items"]
    by_holder = {item["participant_name"]: item for item in items}
    assert set(by_holder) == {AUTHORS[0], AUTHORS[3]}
    assert by_holder[AUTHORS[0]]["authors"] == list(AUTHORS[:3])
    assert by_holder[AUTHORS[0]]["team_name"] == "Equipe Verde"
    assert by_holder[AUTHORS[0]]["work_title"] == "Horta vertical automatizada"
    assert by_holder[AUTHORS[0]]["workload_hours"] is None

    # Any author finds the certificate, in the panel and in the public search.
    found = client.get("/certificates", params={"name": AUTHORS[2]}).json()["items"]
    assert [item["participant_name"] for item in found] == [AUTHORS[0]]

    code = by_holder[AUTHORS[0]]["unique_code"]
    # The stored PDF is named after the work, not its first author.
    assert db.get_by_code(code)["original_filename"] == (
        f"Horta vertical automatizada - Mostra de Iniciação Científica 2026 - {code}.pdf"
    )
    public = _public_client(monkeypatch)
    verify = public.get(f"/public/verify/{code}").json()
    assert verify["certificate"]["per_work"] is True
    assert verify["certificate"]["authors"] == list(AUTHORS[:3])
    assert verify["certificate"]["work_title"] == "Horta vertical automatizada"

    page = public.get(f"/validar/{code}").text
    assert "Autores" in page
    for name in AUTHORS[:3]:
        assert name in page
    assert "Título do trabalho" in page and "Equipe Verde" in page

    search = public.post("/public/search", json={"nome": AUTHORS[2]}).json()
    assert search["total"] == 1
    assert search["items"][0]["authors"] == [
        "P••• E••• A•••",
        "P••• E••• B•••",
        "P••• E••• G•••",
    ]
    assert AUTHORS[2] not in str(search)


def test_blank_workload_used_by_the_body_only_warns(client, auth_headers):
    mostra = _create(client, "Mostra", MOSTRA, MOSTRA_ELEMENTS)
    assert client.post(f"/templates/versions/{mostra['id']}/activate").status_code == 200

    preview = client.post(
        "/certificates/validate-spreadsheet",
        files=_files([_work(AUTHORS[0], carga_horaria=8), _work(AUTHORS[1])]),
        data={
            "texto_padrao": "participou com {{carga_horaria}} horas.",
            "certificate_type": MOSTRA,
        },
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["empty_fields"] == [
        {
            "key": "carga_horaria",
            "label": "Carga horária",
            "row_count": 1,
            "rows": [{"row_number": 3, "nome": AUTHORS[1]}],
            "checked_rows": 2,
            "truncated": False,
        }
    ]


def test_spreadsheet_template_offers_author_columns_for_per_work_types(client, auth_headers):
    from openpyxl import load_workbook

    mostra = client.get("/certificates/spreadsheet-template", params={"certificate_type": MOSTRA})
    assert mostra.status_code == 200
    headers = [cell.value for cell in load_workbook(BytesIO(mostra.content)).active[1]]
    assert headers[:8] == [f"autor{index}" for index in range(1, 9)]
    assert {"evento", "titulo_trabalho", "equipe", "carga_horaria"} <= set(headers)
    assert "nome" not in headers

    generic = client.get("/certificates/spreadsheet-template")
    generic_headers = [cell.value for cell in load_workbook(BytesIO(generic.content)).active[1]]
    assert generic_headers[:2] == ["nome", "carga_horaria"]
    assert {"equipe", "titulo_trabalho"} <= set(generic_headers)


def test_ordinary_certificates_keep_their_holder_view(client, auth_headers):
    """A certificate without authors still shows the participant, not an empty list."""
    from database.repositories import authors_from_json, authors_search_text

    assert authors_from_json(None) == []
    assert authors_from_json("not json") == []
    assert authors_search_text(None) is None
    assert authors_search_text('["Pessoa Exemplo Alfa", "  "]') == "pessoa exemplo alfa"
